from decimal import Decimal
from io import BytesIO

from rest_framework.test import APITestCase

from apps.accounts.models import User
from apps.clients.models import Client
from apps.inventory.models import Product
from apps.sales.models import Contract, ContractDocument, ContractItem


def _make_docx_bytes(paragraphs):
    """Testlar uchun HAQIQIY (docxtpl o'qiy oladigan) minimal `.docx`."""
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


class ContractDocumentTests(APITestCase):
    """13-§1: shartnoma matni — bugalter yuklaydi va saytda tahrirlaydi."""

    def setUp(self):
        self.sales = User.objects.create_user('sal', password='p', role=User.Role.SALES)
        self.sales2 = User.objects.create_user('sal2', password='p', role=User.Role.SALES)
        self.bugalter = User.objects.create_user('bug', password='p', role=User.Role.BUGALTER)
        self.admin = User.objects.create_user('adm', password='p', role=User.Role.ADMIN)
        self.engineer = User.objects.create_user('eng', password='p', role=User.Role.ENGINEER)
        self.supplier = User.objects.create_user('buy', password='p', role=User.Role.SUPPLIER)
        self.mijoz = Client.objects.create(
            type=Client.Type.INDIVIDUAL, full_name='Ali Valiyev',
            passport='AA1112223', jshshir='11112222333344', phone='+998900000001',
        )
        self.product = Product.objects.create(
            sku='HP-880', name='HP 880', sale_price=Decimal('12000000'),
        )
        self.contract = Contract.objects.create(
            client=self.mijoz, total_amount=Decimal('13440000'),
            created_by=self.sales, status=Contract.Status.PENDING_BUGALTER,
        )
        ContractItem.objects.create(
            contract=self.contract, product=self.product, quantity=1,
            unit_price=Decimal('12000000'),
        )

    def test_get_lazily_creates_document(self):
        self.client.force_authenticate(self.bugalter)
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['body'], '')
        self.assertTrue(response.data['can_edit'])
        self.assertTrue(ContractDocument.objects.filter(contract=self.contract).exists())

    def test_engineer_and_supplier_have_no_access(self):
        for user in (self.engineer, self.supplier):
            self.client.force_authenticate(user)
            response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
            self.assertEqual(response.status_code, 403, (user.username, response.data))

    def test_sales_owner_can_read(self):
        self.client.force_authenticate(self.sales)
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(response.data['can_edit'])

    def test_other_sales_cannot_read(self):
        """Egasi bo'lmagan sales — get_queryset uni umuman ko'rmaydi (404)."""
        self.client.force_authenticate(self.sales2)
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertEqual(response.status_code, 404)

    def test_admin_can_read(self):
        """13-§1: tahrir hozircha FAQAT bugalterda — admin ham yo'q."""
        self.client.force_authenticate(self.admin)
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(response.data['can_edit'])

    def test_put_document_no_longer_exists(self):
        """QOLGAN-ISHLAR-2 §4: HTML tahriri (`PUT`) butunlay olib tashlandi —
        Didoxga `source_file` ketadi, `body`ni alohida yozish ikkinchi
        (eski) manba yaratardi. Admin bilan — permission emas, aynan
        metod yo'qligini tekshiramiz (405)."""
        self.client.force_authenticate(self.admin)
        response = self.client.put(
            f'/api/contracts/{self.contract.id}/document/',
            {'body': 'x'}, format='json',
        )
        self.assertEqual(response.status_code, 405)

    def test_bugalter_saves_new_version_on_each_upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        r1 = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': SimpleUploadedFile(
                'v1.docx', _make_docx_bytes(['Birinchi']),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )}, format='multipart',
        )
        self.assertEqual(r1.status_code, 200, r1.data)
        r2 = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': SimpleUploadedFile(
                'v2.docx', _make_docx_bytes(['Ikkinchi']),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )}, format='multipart',
        )
        self.assertEqual(r2.status_code, 200, r2.data)

        versions = self.client.get(
            f'/api/contracts/{self.contract.id}/document/versions/',
        )
        self.assertEqual(len(versions.data), 2)
        self.assertIn('Ikkinchi', versions.data[0]['body'])  # -created_at

    def test_upload_locked_after_didox_sent(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        Contract.objects.filter(pk=self.contract.pk).update(
            status=Contract.Status.PENDING_DIDOX,
        )
        self.client.force_authenticate(self.bugalter)
        response = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': SimpleUploadedFile(
                'x.docx', _make_docx_bytes(['x']),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )}, format='multipart',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Didoxga', str(response.data['detail']))

    def test_upload_rejects_non_docx(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        response = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': SimpleUploadedFile('shartnoma.pdf', b'%PDF-1.4', content_type='application/pdf')},
            format='multipart',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('file', response.data)

    def test_upload_docx_converts_to_html_and_keeps_source_file(self):
        """13-§1 bosqich 3: `.docx` -> mammoth -> body; asl fayl saqlanadi."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        fake_docx = SimpleUploadedFile(
            'shartnoma.docx', _make_docx_bytes(['Yuklangan matn']),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': fake_docx}, format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn('Yuklangan matn', response.data['body'])
        self.assertTrue(response.data['source_file'])
        self.assertTrue(response.data['source_file_name'].startswith('shartnoma'))
        self.assertTrue(response.data['source_file_name'].endswith('.docx'))
        self.assertIsNotNone(response.data['source_uploaded_at'])

        document = ContractDocument.objects.get(contract=self.contract)
        self.assertIn('Yuklangan matn', document.body)
        self.assertTrue(document.source_file.name)
        self.assertEqual(document.versions.count(), 1)
        self.assertEqual(document.docx_version, 1)

    def test_upload_fills_placeholders_statically(self):
        """15-§A: shablondagi `{{ }}` yuklashda TO'LADI — fayl Didoxga shu holicha ketadi."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        template = SimpleUploadedFile(
            'shablon.docx',
            _make_docx_bytes(['Shartnoma: {{ contract.number }}', 'Mijoz: {{ client.name }}']),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': template}, format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn(self.contract.number, response.data['body'])
        self.assertIn('Ali Valiyev', response.data['body'])
        self.assertNotIn('{{', response.data['body'])

    def test_upload_template_can_reference_items_and_totals(self):
        """QOLGAN-ISHLAR-2 §1: shablon endi `items`/`items_total`/`vat_total`ni ko'radi."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        template = SimpleUploadedFile(
            'shablon.docx',
            _make_docx_bytes([
                'Band: {{ items.0.name }} ({{ items.0.sku }}) x{{ items.0.quantity }} = {{ items.0.total }}',
                'Jami (QQSsiz): {{ items_total }}, QQS: {{ vat_total }}',
            ]),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': template}, format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn(self.product.name, response.data['body'])
        self.assertIn(self.product.sku, response.data['body'])
        self.assertIn('13440000', response.data['body'])  # items.0.total (QQS bilan)
        self.assertIn('12000000', response.data['body'])  # items_total
        self.assertIn('1440000', response.data['body'])  # vat_total
        self.assertNotIn('{{', response.data['body'])

    def test_source_file_downloadable_via_document_get(self):
        """15-§3: yuklangan asl `.docx` GET /document/ orqali qaytib olinadi."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        fake_docx = SimpleUploadedFile(
            'referat.docx', _make_docx_bytes(['Referat']),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': fake_docx}, format='multipart',
        )

        self.client.force_authenticate(self.admin)
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn('referat', response.data['source_file'])
        self.assertTrue(response.data['source_file_name'].startswith('referat'))
        self.assertTrue(response.data['source_file_name'].endswith('.docx'))
        self.assertIsNotNone(response.data['source_uploaded_at'])

    def test_source_file_null_before_upload(self):
        """15-§3 regressiya: hech qachon yuklanmagan bo'lsa — hammasi `null`."""
        self.client.force_authenticate(self.bugalter)
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIsNone(response.data['source_file'])
        self.assertIsNone(response.data['source_file_name'])
        self.assertIsNone(response.data['source_uploaded_at'])
        self.assertFalse(response.data['is_stale'])

    def test_is_stale_after_total_changes_post_upload(self):
        """QOLGAN-ISHLAR-2 §3: fayl yuklangandan keyin summa o'zgarsa — `is_stale`."""
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        response = self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': SimpleUploadedFile(
                'shartnoma.docx', _make_docx_bytes(['Matn']),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )}, format='multipart',
        )
        self.assertEqual(response.status_code, 200, response.data)
        self.assertFalse(response.data['is_stale'])

        Contract.objects.filter(pk=self.contract.pk).update(
            total_amount=Decimal('20000000'),
        )
        response = self.client.get(f'/api/contracts/{self.contract.id}/document/')
        self.assertTrue(response.data['is_stale'])


class CollaboraWopiTests(APITestCase):
    """15-§A: Collabora Online (WOPI) — `.docx` brauzerda to'g'ridan-to'g'ri tahrirlanadi."""

    def setUp(self):
        self.sales = User.objects.create_user('sal', password='p', role=User.Role.SALES)
        self.bugalter = User.objects.create_user('bug', password='p', role=User.Role.BUGALTER)
        self.admin = User.objects.create_user('adm', password='p', role=User.Role.ADMIN)
        self.mijoz = Client.objects.create(
            type=Client.Type.INDIVIDUAL, full_name='Ali Valiyev',
            passport='AA1112223', jshshir='11112222333344', phone='+998900000001',
        )
        self.product = Product.objects.create(
            sku='HP-880', name='HP 880', sale_price=Decimal('12000000'),
        )
        self.contract = Contract.objects.create(
            client=self.mijoz, total_amount=Decimal('13440000'),
            created_by=self.sales, status=Contract.Status.PENDING_BUGALTER,
        )
        ContractItem.objects.create(
            contract=self.contract, product=self.product, quantity=1,
            unit_price=Decimal('12000000'),
        )

    def _upload(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        self.client.force_authenticate(self.bugalter)
        return self.client.post(
            f'/api/contracts/{self.contract.id}/document/upload/',
            {'file': SimpleUploadedFile(
                'shartnoma.docx', _make_docx_bytes(['Matn']),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            )}, format='multipart',
        )

    def test_edit_session_requires_uploaded_docx(self):
        self.client.force_authenticate(self.bugalter)
        response = self.client.post(f'/api/contracts/{self.contract.id}/document/edit-session/')
        self.assertEqual(response.status_code, 400)

    def test_edit_session_only_bugalter(self):
        self._upload()
        for user in (self.sales, self.admin):
            self.client.force_authenticate(user)
            response = self.client.post(f'/api/contracts/{self.contract.id}/document/edit-session/')
            self.assertEqual(response.status_code, 403, (user.username, response.data))

    def test_edit_session_returns_collabora_url_with_token(self):
        self._upload()
        self.client.force_authenticate(self.bugalter)
        response = self.client.post(f'/api/contracts/{self.contract.id}/document/edit-session/')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertIn('WOPISrc=', response.data['edit_url'])
        self.assertIn('access_token=', response.data['edit_url'])
        self.assertIn('/wopi/files/', response.data['wopi_src'])

    def _token(self):
        from apps.sales.services import create_wopi_token, get_or_create_contract_document

        document = get_or_create_contract_document(self.contract)
        return document, create_wopi_token(document, self.bugalter)

    def test_check_file_info_requires_valid_token(self):
        self._upload()
        document, token = self._token()
        response = self.client.get(f'/api/wopi/files/{document.pk}?access_token=bad')
        self.assertEqual(response.status_code, 401)

        response = self.client.get(f'/api/wopi/files/{document.pk}?access_token={token}')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data['Version'], '1')
        self.assertTrue(response.data['BaseFileName'].endswith('.docx'))
        self.assertTrue(response.data['UserCanWrite'])

    def test_get_file_returns_raw_docx_bytes(self):
        self._upload()
        document, token = self._token()
        response = self.client.get(f'/api/wopi/files/{document.pk}/contents?access_token={token}')
        self.assertEqual(response.status_code, 200)
        document.refresh_from_db()
        document.source_file.open('rb')
        self.assertEqual(response.content, document.source_file.read())
        document.source_file.close()

    def test_lock_conflict_returns_409_with_current_lock(self):
        self._upload()
        document, token = self._token()
        url = f'/api/wopi/files/{document.pk}?access_token={token}'
        r1 = self.client.post(url, HTTP_X_WOPI_OVERRIDE='LOCK', HTTP_X_WOPI_LOCK='lock-a')
        self.assertEqual(r1.status_code, 200)

        r2 = self.client.post(url, HTTP_X_WOPI_OVERRIDE='LOCK', HTTP_X_WOPI_LOCK='lock-b')
        self.assertEqual(r2.status_code, 409)
        self.assertEqual(r2['X-WOPI-Lock'], 'lock-a')

        r3 = self.client.post(url, HTTP_X_WOPI_OVERRIDE='UNLOCK', HTTP_X_WOPI_LOCK='lock-a')
        self.assertEqual(r3.status_code, 200)

        r4 = self.client.post(url, HTTP_X_WOPI_OVERRIDE='LOCK', HTTP_X_WOPI_LOCK='lock-b')
        self.assertEqual(r4.status_code, 200)

    def test_put_file_saves_new_docx_and_bumps_version(self):
        self._upload()
        document, token = self._token()
        url = f'/api/wopi/files/{document.pk}/contents?access_token={token}'
        new_bytes = _make_docx_bytes(['Collaboradan saqlangan matn'])
        response = self.client.post(url, data=new_bytes, content_type='application/octet-stream')
        self.assertEqual(response.status_code, 200)

        document.refresh_from_db()
        self.assertEqual(document.docx_version, 2)
        self.assertIn('Collaboradan saqlangan matn', document.body)
        self.assertEqual(document.versions.count(), 2)

    def test_put_file_blocked_by_foreign_lock(self):
        self._upload()
        document, token = self._token()
        self.client.post(
            f'/api/wopi/files/{document.pk}?access_token={token}',
            HTTP_X_WOPI_OVERRIDE='LOCK', HTTP_X_WOPI_LOCK='someone-elses-lock',
        )
        response = self.client.post(
            f'/api/wopi/files/{document.pk}/contents?access_token={token}',
            data=_make_docx_bytes(['x']), content_type='application/octet-stream',
            HTTP_X_WOPI_LOCK='different-lock',
        )
        self.assertEqual(response.status_code, 409)
